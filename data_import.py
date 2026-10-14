"""Validated, repeatable CSV/JSON imports for sourced healthcare directory records."""

from __future__ import annotations

import csv
from datetime import date, datetime
from hashlib import sha256
import io
import json
import re
import sqlite3
from collections.abc import Iterable, Mapping
from urllib.parse import urlsplit

from openpyxl import load_workbook

DIRECTORY_CATEGORIES = {"cancer", "dialysis", "hospital", "blood_bank", "doctor"}
VERIFICATION_STATUSES = {"pending", "verified", "unverified"}
IMPORT_FIELDS = {
    "name",
    "category",
    "address",
    "city",
    "area",
    "phone",
    "website",
    "ownership_type",
    "specialty",
    "source_url",
    "last_verified",
    "verification_status",
    "scheme_name",
    "empanelment_verification_status",
    "scheme_source_url",
    "scheme_last_verified",
    "scheme_specialties",
}
TEXT_LIMITS = {
    "name": 200,
    "address": 400,
    "city": 100,
    "area": 150,
    "phone": 100,
    "website": 500,
    "ownership_type": 20,
    "specialty": 300,
    "source_url": 500,
    "scheme_source_url": 500,
    "scheme_specialties": 500,
}
EXCEL_COLUMNS = {
    "Facility Name": "name",
    "Category": "source_category",
    "Address": "address",
    "Phone": "phone",
    "Services": "services",
    "Latitude": "latitude",
    "Longitude": "longitude",
    "Blood Group / Stock": "blood_stock",
    "Hours": "opening_hours",
    "Verification Status": "source_verification_status",
    "Source URL": "source_url",
    "Last Checked": "last_checked",
    "Notes": "notes",
}
OPTIONAL_EXCEL_COLUMNS = {
    "Facility Type": "ownership_type",
    "Scheme Name": "scheme_name",
    "Empanelment Verification Status": "empanelment_verification_status",
    "Scheme Source URL": "scheme_source_url",
    "Scheme Last Verified": "scheme_last_verified",
    "Scheme Specialties": "scheme_specialties",
}
AYUSHMAN_SCHEMES = {"PM-JAY", "MJPJAY"}


class DirectoryImportError(ValueError):
    """Raised when an import file or record is invalid."""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _category(value: str) -> str:
    normalized = value.casefold()
    if "directory search" in normalized or "lookup" in normalized:
        return ""
    if "cancer" in normalized or "oncolog" in normalized:
        return "cancer"
    if "dialysis" in normalized or "kidney" in normalized:
        return "dialysis"
    if "blood bank" in normalized or "blood centre" in normalized or "blood center" in normalized:
        return "blood_bank"
    if "hospital" in normalized or "medical college" in normalized or "dental college" in normalized:
        return "hospital"
    return ""


def _verification_status(source_status: str) -> str:
    normalized = source_status.casefold()
    if "pending" in normalized or "unverified" in normalized or "verify" in normalized or "confirm" in normalized or "recheck" in normalized or "recommend" in normalized:
        return "pending"
    if normalized.strip() in {"verified", "officially verified"}:
        return "verified"
    return "unverified"


def parse_healthcare_workbook(content: bytes) -> dict[str, object]:
    """Read facilities, statistics, and submission-template headers from the supplied workbook."""
    if len(content) > 8 * 1024 * 1024:
        raise DirectoryImportError("Excel files must be 8 MB or smaller.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise DirectoryImportError("The uploaded Excel workbook is invalid or unreadable.") from error

    required_sheets = {"Facility Directory", "Data Analysis", "Add New Records Template"}
    missing_sheets = required_sheets - set(workbook.sheetnames)
    if missing_sheets:
        raise DirectoryImportError(
            "Workbook is missing required sheets: " + ", ".join(sorted(missing_sheets))
        )
    facility_sheet = workbook["Facility Directory"]
    rows = facility_sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise DirectoryImportError("Facility Directory has no header row.")
    header_positions = {_cell_text(value): index for index, value in enumerate(header) if value is not None}
    missing_columns = set(EXCEL_COLUMNS) - set(header_positions)
    if missing_columns:
        raise DirectoryImportError(
            "Facility Directory is missing required columns: "
            + ", ".join(sorted(missing_columns))
        )

    records: list[dict[str, object]] = []
    skipped = 0
    row_errors: list[str] = []
    for sheet_row_number, cells in enumerate(rows, start=2):
        values = {
            field: _cell_text(cells[position] if position < len(cells) else None)
            for heading, field in EXCEL_COLUMNS.items()
            for position in [header_positions[heading]]
        }
        values.update({
            field: _cell_text(cells[header_positions[heading]])
            for heading, field in OPTIONAL_EXCEL_COLUMNS.items()
            if heading in header_positions
        })
        if not values["name"] and not values["source_category"]:
            continue
        category = _category(values["source_category"])
        if not category:
            skipped += 1
            continue
        try:
            source_url = _https_url(values["source_url"], "source_url", required=True)
            phone = values["phone"]
            if phone and (
                not re.fullmatch(r"[\w+() .;:/,-]{1,100}", phone, re.UNICODE)
            ):
                raise DirectoryImportError("phone field contains unsupported characters.")
            latitude = float(values["latitude"]) if values["latitude"] else None
            longitude = float(values["longitude"]) if values["longitude"] else None
            if latitude is not None and not -90 <= latitude <= 90:
                raise DirectoryImportError("latitude is outside the valid range.")
            if longitude is not None and not -180 <= longitude <= 180:
                raise DirectoryImportError("longitude is outside the valid range.")
            checked = values["last_checked"]
            if checked:
                checked = date.fromisoformat(checked[:10]).isoformat()
        except (ValueError, DirectoryImportError) as error:
            row_errors.append(f"Row {sheet_row_number}: {error}")
            skipped += 1
            continue

        address = values["address"]
        name = values["name"]
        if not name or not address:
            row_errors.append(f"Row {sheet_row_number}: facility name or address is missing.")
        city = "Nagpur"
        source_key = "|".join(
            " ".join(part.casefold().split())
            for part in (name, category, city, address)
        )
        facility_id = "xlsx-" + sha256(source_key.encode("utf-8")).hexdigest()[:24]
        records.append({
            "facility_id": facility_id,
            "name": name,
            "category": category,
            "source_category": values["source_category"],
            "address": address,
            "city": city,
            "area": "",
            "phone": phone,
            "website": "",
            "ownership_type": values.get("ownership_type", ""),
            "specialty": "",
            "services": values["services"],
            "latitude": latitude,
            "longitude": longitude,
            "blood_stock": values["blood_stock"],
            "opening_hours": values["opening_hours"],
            "source_url": source_url,
            "last_verified": None,
            "last_checked": checked or None,
            "verification_status": _verification_status(values["source_verification_status"]),
            "source_verification_status": values["source_verification_status"],
            "notes": values["notes"],
            "scheme_name": values.get("scheme_name", ""),
            "empanelment_verification_status": values.get(
                "empanelment_verification_status", ""
            ),
            "scheme_source_url": values.get("scheme_source_url", ""),
            "scheme_last_verified": values.get("scheme_last_verified", ""),
            "scheme_specialties": values.get("scheme_specialties", ""),
        })

    analysis_sheet = workbook["Data Analysis"]
    analysis_rows = list(analysis_sheet.iter_rows(values_only=True))
    template_rows = workbook["Add New Records Template"].iter_rows(values_only=True)
    template_headers = [_cell_text(value) for value in next(template_rows, ()) if value is not None]
    analysis = [
        {"metric": _cell_text(row[0]), "value": _cell_text(row[1]), "interpretation": _cell_text(row[2])}
        for row in analysis_rows[1:]
        if row and len(row) >= 2 and row[0] is not None
    ]
    workbook.close()
    return {
        "records": records,
        "records_read": len(records) + skipped,
        "skipped": skipped,
        "errors": row_errors,
        "analysis": analysis,
        "template_headers": template_headers,
    }


def _https_url(value: str, field: str, required: bool = False) -> str:
    if not value and not required:
        return ""
    parsed = urlsplit(value)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise DirectoryImportError(f"{field} must be a valid HTTPS URL.")
    return value


def parse_directory_file(filename: str, content: bytes) -> list[dict[str, str]]:
    """Parse UTF-8 CSV or JSON records without fetching any remote source."""
    if len(content) > 8 * 1024 * 1024:
        raise DirectoryImportError("Import files must be 8 MB or smaller.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DirectoryImportError("Import files must use UTF-8 encoding.") from error

    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        rows = list(csv.DictReader(io.StringIO(text)))
    elif suffix == "json":
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as error:
            raise DirectoryImportError("The JSON file is malformed.") from error
        if isinstance(payload, dict):
            payload = payload.get("records", payload.get("listings"))
        if not isinstance(payload, list) or any(not isinstance(row, dict) for row in payload):
            raise DirectoryImportError("JSON must contain a list of records or a records/listings array.")
        rows = payload
    else:
        raise DirectoryImportError("Choose a .csv or .json file.")

    if not rows:
        raise DirectoryImportError("The import file contains no records.")
    return [{str(key): "" if value is None else str(value) for key, value in row.items()} for row in rows]


def _normalized_key(name: str, category: str, city: str) -> tuple[str, str, str]:
    normalize = lambda value: " ".join(value.casefold().split())
    return normalize(name), category, normalize(city)


def _validated_row(raw: Mapping[str, object], index: int) -> dict[str, str]:
    row = {field: str(raw.get(field, "") or "").strip() for field in IMPORT_FIELDS}
    prefix = f"Record {index}: "
    if not row["name"] or not row["city"] or not row["source_url"]:
        raise DirectoryImportError(prefix + "name, city, and source_url are required.")
    if row["category"] not in DIRECTORY_CATEGORIES:
        raise DirectoryImportError(prefix + "category must be cancer, dialysis, hospital, blood_bank, or doctor.")
    row["ownership_type"] = row["ownership_type"].casefold()
    if row["ownership_type"] and row["ownership_type"] not in {
        "government", "private", "unknown"
    }:
        raise DirectoryImportError(prefix + "ownership_type must be government, private, or unknown.")
    if row["verification_status"] and row["verification_status"] not in VERIFICATION_STATUSES:
        raise DirectoryImportError(prefix + "verification_status must be pending, verified, or unverified.")
    for field, maximum in TEXT_LIMITS.items():
        if len(row[field]) > maximum:
            raise DirectoryImportError(prefix + f"{field} must be {maximum} characters or fewer.")
    _https_url(row["source_url"], "source_url", required=True)
    _https_url(row["website"], "website")
    if row["phone"] and (
        not re.fullmatch(r"[\w+() .;:/,-]{1,100}", row["phone"], re.UNICODE)
    ):
        raise DirectoryImportError(prefix + "phone field contains unsupported characters.")
    if row["last_verified"]:
        try:
            date.fromisoformat(row["last_verified"])
        except ValueError as error:
            raise DirectoryImportError(prefix + "last_verified must use YYYY-MM-DD.") from error
    if row["verification_status"] == "verified" and not row["last_verified"]:
        raise DirectoryImportError(prefix + "verified records require last_verified.")
    result: dict[str, object] = dict(row)
    for field, maximum in (
        ("services", 1000),
        ("blood_stock", 300),
        ("opening_hours", 300),
        ("notes", 2000),
        ("source_category", 200),
        ("source_verification_status", 300),
        ("last_checked", 30),
    ):
        value = str(raw.get(field, "") or "").strip()
        if len(value) > maximum:
            raise DirectoryImportError(prefix + f"{field} must be {maximum} characters or fewer.")
        result[field] = value
    result["facility_id"] = str(raw.get("facility_id", "") or "").strip()[:100]
    result["latitude"] = raw.get("latitude")
    result["longitude"] = raw.get("longitude")
    for field, minimum, maximum in (
        ("latitude", -90, 90),
        ("longitude", -180, 180),
    ):
        value = result[field]
        if value in (None, ""):
            result[field] = None
            continue
        try:
            number = float(value)
        except (ValueError, TypeError) as error:
            raise DirectoryImportError(prefix + f"{field} must be numeric.") from error
        if not minimum <= number <= maximum:
            raise DirectoryImportError(prefix + f"{field} is outside the valid range.")
        result[field] = number
    for field in ("last_checked",):
        value = str(result[field] or "")
        if value:
            try:
                result[field] = date.fromisoformat(value[:10]).isoformat()
            except ValueError as error:
                raise DirectoryImportError(prefix + f"{field} must use YYYY-MM-DD.") from error
    if result["blood_stock"]:
        result["blood_stock"] = str(result["blood_stock"])[:300]
    result["scheme_name"] = str(raw.get("scheme_name", "") or "").strip()
    result["empanelment_verification_status"] = str(
        raw.get("empanelment_verification_status", "") or ""
    ).strip().lower()
    result["scheme_source_url"] = str(raw.get("scheme_source_url", "") or "").strip()
    result["scheme_last_verified"] = str(raw.get("scheme_last_verified", "") or "").strip()
    result["scheme_specialties"] = str(raw.get("scheme_specialties", "") or "").strip()
    if any(result[field] for field in (
        "empanelment_verification_status",
        "scheme_source_url",
        "scheme_last_verified",
        "scheme_specialties",
    )) and not result["scheme_name"]:
        raise DirectoryImportError(prefix + "scheme_name is required for scheme details.")
    if result["scheme_name"]:
        if result["scheme_name"] not in AYUSHMAN_SCHEMES:
            raise DirectoryImportError(prefix + "scheme_name must be PM-JAY or MJPJAY.")
        if result["empanelment_verification_status"] not in VERIFICATION_STATUSES:
            raise DirectoryImportError(
                prefix + "empanelment_verification_status must be pending, verified, or unverified."
            )
        if not result["scheme_source_url"]:
            raise DirectoryImportError(prefix + "scheme_source_url is required for scheme details.")
        _https_url(result["scheme_source_url"], "scheme_source_url", required=True)
        scheme_host = urlsplit(result["scheme_source_url"]).hostname or ""
        scheme_host = scheme_host.casefold().rstrip(".")
        if not (scheme_host.endswith(".gov.in") or scheme_host.endswith(".nic.in")):
            raise DirectoryImportError(
                prefix + "scheme_source_url must use an official .gov.in or .nic.in domain."
            )
        if len(result["scheme_source_url"]) > 500:
            raise DirectoryImportError(prefix + "scheme_source_url must be 500 characters or fewer.")
        if len(result["scheme_specialties"]) > 500:
            raise DirectoryImportError(prefix + "scheme_specialties must be 500 characters or fewer.")
        if result["scheme_last_verified"]:
            try:
                result["scheme_last_verified"] = date.fromisoformat(
                    result["scheme_last_verified"]
                ).isoformat()
            except ValueError as error:
                raise DirectoryImportError(
                    prefix + "scheme_last_verified must use YYYY-MM-DD."
                ) from error
        if result["empanelment_verification_status"] == "verified":
            if not result["scheme_last_verified"]:
                raise DirectoryImportError(
                    prefix + "verified scheme records require scheme_last_verified."
                )
            if not result["scheme_specialties"]:
                raise DirectoryImportError(
                    prefix + "verified scheme records require scheme_specialties."
                )
    elif result["scheme_name"] == "":
        result["empanelment_verification_status"] = ""
        result["scheme_source_url"] = ""
        result["scheme_last_verified"] = ""
        result["scheme_specialties"] = ""
    return result


def _identity_key(row: Mapping[str, object]) -> tuple[str, str, str, str]:
    normalize = lambda value: " ".join(str(value or "").casefold().split())
    return (
        normalize(row.get("name")),
        str(row.get("category", "")),
        normalize(row.get("city")),
        normalize(row.get("address")),
    )


def _scheme_row_changed(
    db: sqlite3.Connection,
    listing_id: int,
    row: Mapping[str, object],
) -> bool:
    scheme_name = str(row.get("scheme_name", "") or "").strip()
    if not scheme_name:
        return False
    existing = db.execute(
        """SELECT empanelment_verification_status, source_url, last_verified, specialties
           FROM ayushman_empanelments WHERE listing_id = ? AND scheme_name = ?""",
        (listing_id, scheme_name),
    ).fetchone()
    expected = (
        str(row.get("empanelment_verification_status", "") or "").strip(),
        str(row.get("scheme_source_url", "") or "").strip(),
        str(row.get("scheme_last_verified", "") or "").strip() or None,
        str(row.get("scheme_specialties", "") or "").strip(),
    )
    return existing is None or tuple(existing) != expected


def preview_directory_import(
    db: sqlite3.Connection,
    records: Iterable[Mapping[str, object]],
) -> dict[str, object]:
    """Classify prospective changes without mutating the directory."""
    validated = [_validated_row(row, index) for index, row in enumerate(records, start=1)]
    if not validated:
        raise DirectoryImportError("The import contains no valid facility records.")
    current = db.execute("SELECT * FROM healthcare_directory").fetchall()
    by_source_id = {
        row["facility_id"]: row
        for row in current
        if row["facility_id"]
    }
    by_identity: dict[tuple[str, str, str, str], sqlite3.Row] = {}
    by_name_location: dict[tuple[str, str, str], list[sqlite3.Row]] = {}
    for existing in current:
        existing_values = dict(existing)
        by_identity[_identity_key(existing_values)] = existing
        key = _identity_key(existing_values)[:3]
        by_name_location.setdefault(key, []).append(existing)

    new_count = updated = unchanged = flagged = 0
    row_matches: list[int | None] = []
    row_flags: list[bool] = []
    duplicate_rows: list[bool] = []
    seen_identities: set[tuple[str, str, str, str]] = set()
    seen_by_name_location: dict[tuple[str, str, str], list[dict[str, object]]] = {}
    compared_fields = (
        "name", "category", "address", "city", "area", "phone", "website",
        "ownership_type",
        "specialty", "services", "latitude", "longitude", "blood_stock",
        "opening_hours", "source_url", "last_checked", "notes",
        "source_verification_status",
    )
    for row in validated:
        identity = _identity_key(row)
        name_location = identity[:3]
        if identity in seen_identities:
            unchanged += 1
            row_matches.append(None)
            row_flags.append(False)
            duplicate_rows.append(True)
            continue

        existing = (
            by_source_id.get(str(row["facility_id"]))
            if row["facility_id"]
            else by_identity.get(identity)
        )
        if existing is not None:
            differences = [
                field for field in compared_fields
                if row.get(field) not in (None, "")
                and str(row.get(field)) != str(existing[field] or "")
            ]
            if differences or _scheme_row_changed(db, existing["id"], row):
                updated += 1
            else:
                unchanged += 1
            row_matches.append(existing["id"])
            row_flags.append(False)
            duplicate_rows.append(False)
            seen_identities.add(identity)
            continue

        candidates = [
            *by_name_location.get(name_location, []),
            *seen_by_name_location.get(name_location, []),
        ]
        if candidates:
            flagged += 1
            row_matches.append(None)
            row_flags.append(True)
        else:
            new_count += 1
            row_matches.append(None)
            row_flags.append(False)
        duplicate_rows.append(False)
        seen_identities.add(identity)
        seen_by_name_location.setdefault(name_location, []).append(row)
    return {
        "records": validated,
        "matches": row_matches,
        "flags": row_flags,
        "duplicate_rows": duplicate_rows,
        "new": new_count,
        "updated": updated,
        "unchanged": unchanged,
        "flagged": flagged,
    }


def import_directory_records(
    db: sqlite3.Connection,
    records: Iterable[Mapping[str, object]],
    submitted_by: int | None = None,
) -> dict[str, int]:
    """Insert new entries and safely merge exact identities; flag branch conflicts."""
    preview = preview_directory_import(db, records)
    rows = preview["records"]
    inserted = updated = unchanged = flagged = 0
    fields = (
        "facility_id", "name", "category", "address", "city", "area", "phone",
        "website", "ownership_type", "specialty", "services", "latitude", "longitude", "blood_stock",
        "opening_hours", "source_url", "last_verified", "last_checked", "notes",
        "source_verification_status", "verification_status", "potentially_duplicate",
        "submitted_by",
    )

    def save_scheme_data(db: sqlite3.Connection, listing_id: int, row: Mapping[str, object]) -> None:
        scheme_name = str(row.get("scheme_name", "") or "").strip()
        if not scheme_name:
            return
        scheme_status = str(row.get("empanelment_verification_status", "") or "").strip()
        scheme_source_url = str(row.get("scheme_source_url", "") or "").strip()
        scheme_last_verified = str(row.get("scheme_last_verified", "") or "").strip() or None
        scheme_specialties = str(row.get("scheme_specialties", "") or "").strip()
        db.execute(
            """INSERT INTO ayushman_empanelments
               (listing_id, scheme_name, empanelment_verification_status,
                source_url, last_verified, specialties)
               VALUES (?, ?, ?, ?, ?, ?)
               ON CONFLICT(listing_id, scheme_name) DO UPDATE SET
                   empanelment_verification_status = excluded.empanelment_verification_status,
                   source_url = excluded.source_url,
                   last_verified = excluded.last_verified,
                   specialties = excluded.specialties,
                   last_updated = CURRENT_TIMESTAMP""",
            (
                listing_id,
                scheme_name,
                scheme_status,
                scheme_source_url,
                scheme_last_verified,
                scheme_specialties,
            ),
        )

    for row, existing_id, potential_duplicate, duplicate_row in zip(
        rows,
        preview["matches"],
        preview["flags"],
        preview["duplicate_rows"],
        strict=True,
    ):
        if duplicate_row:
            unchanged += 1
            continue
        if potential_duplicate:
            flagged += 1
        if existing_id is None:
            values = []
            for field in fields[:-3]:
                value = row.get(field, "")
                if field == "ownership_type" and not value:
                    value = "unknown"
                if field in ("facility_id", "last_verified", "last_checked"):
                    value = value or None
                values.append(value)
            values.extend((
                "pending" if potential_duplicate else row["verification_status"] or "unverified",
                int(potential_duplicate),
                submitted_by,
            ))
            db.execute(
                f"INSERT INTO healthcare_directory ({', '.join(fields)}) "
                f"VALUES ({', '.join('?' for _ in fields)})",
                values,
            )
            save_scheme_data(db, db.execute("SELECT last_insert_rowid()").fetchone()[0], row)
            inserted += 1
            continue

        existing = db.execute(
            "SELECT * FROM healthcare_directory WHERE id = ?", (existing_id,)
        ).fetchone()
        update_fields = (
            "name", "category", "address", "city", "area", "phone", "website",
            "ownership_type", "specialty", "services", "latitude", "longitude", "blood_stock",
            "opening_hours", "source_url", "last_checked", "notes",
            "source_verification_status",
        )
        changed = any(
            row.get(field) not in (None, "")
            and str(row.get(field)) != str(existing[field] or "")
            for field in update_fields
        )
        scheme_changed = _scheme_row_changed(db, existing_id, row)
        if not changed and not scheme_changed:
            unchanged += 1
            continue
        if not changed:
            save_scheme_data(db, existing_id, row)
            updated += 1
            continue
        assignments = [
            f"{field} = COALESCE(NULLIF(?, ''), {field})"
            for field in update_fields
        ]
        values = [row.get(field, "") for field in update_fields]
        assignments.extend(
            ("last_updated = CURRENT_TIMESTAMP", "verification_status = ?")
        )
        values.append(
            row["verification_status"]
            if row["verification_status"] == "verified"
            else "pending"
        )
        if row.get("last_verified"):
            assignments.append("last_verified = ?")
            values.append(row["last_verified"])
        values.append(existing_id)
        db.execute(
            f"UPDATE healthcare_directory SET {', '.join(assignments)} WHERE id = ?",
            values,
        )
        save_scheme_data(db, existing_id, row)
        updated += 1

    return {
        "inserted": inserted,
        "updated": updated,
        "unchanged": unchanged,
        "flagged": flagged,
        "skipped": 0,
    }
